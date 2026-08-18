from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent

NAVY = "123B5D"
TEAL = "1F7A8C"
LIGHT_TEAL = "EAF6F7"
PALE = "F7FAFC"
GRID = "D7E2E9"
MUTED = "64748B"
TEXT = "243447"
GOLD = "F5E4AD"
GOLD_TEXT = "7A5A00"
GREEN = "EAF7F0"
GREEN_TEXT = "166534"
RED = "FDECEC"
RED_TEXT = "9D2929"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size=11, color=TEXT, bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_para(paragraph, before=0, after=6, line=1.25, align=None, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    if keep:
        fmt.keep_with_next = True


def add_text(doc, text, size=11, color=TEXT, bold=False, italic=False, after=6, before=0, align=None):
    p = doc.add_paragraph()
    set_para(p, before, after, 1.25, align)
    r = p.add_run(text)
    set_run_font(r, size, color, bold, italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    if level == 1:
        set_run_font(r, 16, TEAL, True)
    elif level == 2:
        set_run_font(r, 13, TEAL, True)
    else:
        set_run_font(r, 12, NAVY, True)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    set_para(p, 0, 4, 1.25)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_para(p, 0, 5, 1.25)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_callout(doc, label, text, fill=LIGHT_TEAL, label_color=TEAL):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_para(p, 0, 0, 1.15)
    a = p.add_run(label + "  ")
    set_run_font(a, 10, label_color, True)
    b = p.add_run(text)
    set_run_font(b, 10, TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_table(doc, headers, rows, widths, header_fill="E8EEF5", font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    for cell, value in zip(hdr.cells, headers):
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        set_para(p, 0, 0, 1.1)
        r = p.add_run(value)
        set_run_font(r, font_size, NAVY, True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            p = cell.paragraphs[0]
            set_para(p, 0, 0, 1.1)
            r = p.add_run(str(value))
            set_run_font(r, font_size, TEXT)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_label_detail(doc, label, value):
    p = doc.add_paragraph()
    set_para(p, 0, 3, 1.15)
    a = p.add_run(label + ": ")
    set_run_font(a, 10.5, NAVY, True)
    b = p.add_run(value)
    set_run_font(b, 10.5, TEXT)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, color, before, after in ((1, 16, TEAL, 18, 10), (2, 13, TEAL, 14, 7), (3, 12, NAVY, 10, 5)):
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_header_footer(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    set_para(p, 0, 0, 1.0)
    left = p.add_run("REMTool  /  PANEL REM")
    set_run_font(left, 8.5, MUTED, True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ""
    set_para(fp, 0, 0, 1.0)
    r = fp.add_run("DESAM Monte Patria  ·  Manual de usuario")
    set_run_font(r, 8, MUTED)
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.add_run(" " * 8)
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    fp._p.append(page)


def add_cover(doc):
    add_text(doc, "MANUAL DE USUARIO", 10, TEAL, True, after=20, before=42)
    p = doc.add_paragraph()
    set_para(p, 0, 8, 1.0)
    r = p.add_run("Panel REM")
    set_run_font(r, 34, NAVY, True)
    p = doc.add_paragraph()
    set_para(p, 0, 18, 1.15)
    r = p.add_run("Revisión, validación y consolidación de planillas REM")
    set_run_font(r, 15, TEAL)
    add_callout(doc, "PROPÓSITO", "Orientar a los usuarios en la carga de planillas, lectura de resultados, revisión de hallazgos, consolidación de Serie A y exportación del informe.")
    add_text(doc, "REMTool · DESAM Monte Patria", 12, NAVY, True, before=30, after=4)
    add_text(doc, "Edición 1.0 · 11 de agosto de 2026", 10, MUTED, after=2)
    add_text(doc, "Documento elaborado a partir de la revisión funcional del PanelREM, su analizador de planillas y los servicios de resumen y consolidación.", 9.5, MUTED, italic=True, after=0)
    doc.add_page_break()


def build():
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc)
    add_cover(doc)

    add_heading(doc, "1. ¿Qué es el Panel REM?", 1)
    add_text(doc, "El Panel REM es el espacio de REMTool para revisar una o varias planillas REM en un mismo período. Analiza cada archivo .xlsm, identifica automáticamente la serie y el establecimiento, informa inconsistencias, calcula un resumen por serie y permite exportar un informe de revisión.")
    add_text(doc, "El panel también permite adjuntar planillas complementarias a una planilla base de Serie A para sumar sus valores y descargar un REM consolidado.")
    add_callout(doc, "FLUJO EN UNA LÍNEA", "Definir período y sector  →  Cargar .xlsm  →  Analizar  →  Revisar estados y detalle  →  Consultar resumen  →  Exportar o consolidar")

    add_heading(doc, "2. Antes de comenzar", 1)
    for item in [
        "Ten disponible la planilla REM en formato .xlsm.",
        "Verifica que el archivo sea una planilla REM vigente y que incluya la hoja NOMBRE.",
        "Selecciona primero el mes, el año y, si corresponde, el sector.",
        "Necesitas conexión con REMTool para analizar archivos y calcular el resumen.",
    ]:
        add_bullet(doc, item)
    add_text(doc, "El sistema lee desde la planilla el CodDEIS, mes, versión y serie, y los contrasta con la configuración del sistema. El establecimiento y su sector se obtienen desde el catálogo de REMTool.")

    add_heading(doc, "3. Cómo ingresar", 1)
    add_text(doc, "En la barra superior selecciona Panel REM. El panel abre con el período anterior al mes actual como selección inicial. Puedes cambiar el mes, el año y el sector en cualquier momento.")
    add_table(doc, ["Elemento", "Qué permite hacer"], [
        ("Mes / Año / Sector", "Define el contexto que se mostrará y contra el cual se validarán las planillas."),
        ("Zona de carga", "Arrastra o selecciona una o varias planillas .xlsm."),
        ("Resumen del período", "Muestra los puntos resumen calculados por serie y categoría."),
        ("Listado de planillas", "Organiza los archivos por sector y serie; permite abrir el detalle."),
        ("Exportar revisión", "Descarga un informe compatible con Word con hallazgos por archivo."),
        ("Limpiar revisiones", "Elimina todas las revisiones guardadas en el navegador actual."),
    ], [2300, 7060])

    add_heading(doc, "4. Procedimiento recomendado", 1)
    add_heading(doc, "Paso 1. Define el contexto de revisión", 2)
    add_text(doc, "En Mes, Año y Sector selecciona el contexto que quieres revisar. Si dejas Todos los sectores, el panel muestra los archivos de todos los sectores disponibles para ese período.")
    add_callout(doc, "RECOMENDACIÓN", "Define el contexto antes de cargar archivos. Así las validaciones de mes, año y establecimiento quedan asociadas a la carga y el resumen se actualiza con mayor claridad.", fill="FFF8E8", label_color=GOLD_TEXT)
    add_heading(doc, "Paso 2. Carga las planillas", 2)
    add_text(doc, "Arrastra una o varias planillas .xlsm a la zona de carga o haz clic para buscarlas. Los archivos se analizan individualmente y la aplicación detecta automáticamente la serie y el establecimiento.")
    add_callout(doc, "IMPORTANTE", "Solo se aceptan archivos con extensión .xlsm. Si seleccionas otro tipo de archivo, la carga queda con error.", fill="FFF8E8", label_color=GOLD_TEXT)
    add_heading(doc, "Paso 3. Revisa el resultado de la carga", 2)
    add_table(doc, ["Estado", "Significado"], [
        ("Pendiente", "Está esperando ser procesado."),
        ("Analizando...", "El sistema está leyendo la planilla."),
        ("Revisado", "Fue analizada y quedó incluida en el resumen."),
        ("Revisado · excluido del resumen", "Fue analizada, pero no cumple alguno de los filtros o datos esperados."),
        ("Error", "El archivo no pudo ser procesado."),
    ], [3000, 6360])
    add_heading(doc, "Paso 4. Consulta el resumen del período", 2)
    add_text(doc, "El resumen se recalcula automáticamente cada vez que cambia la selección o se incorpora una planilla. Los resultados se muestran por Serie y por Categoría mediante tarjetas con nombre del punto resumen, valor total calculado y cantidad de planillas consideradas.")
    add_text(doc, "Solo se consideran las planillas marcadas como coincidentes. Las planillas excluidas siguen visibles para facilitar la revisión, pero no aportan al total.")
    add_heading(doc, "Paso 5. Revisa cada planilla", 2)
    add_text(doc, "En Planillas que cumplen los filtros, abre el grupo de la serie y selecciona el archivo. Se abrirá un panel lateral con establecimiento, archivo, fecha de carga, serie, versión, mes, año, CodDEIS, sector, inconsistencias, errores, advertencias y planillas complementarias.")
    add_text(doc, "Si no hay hallazgos, el detalle indica que la planilla fue analizada sin errores ni advertencias.")
    add_heading(doc, "Paso 6. Corrige y vuelve a revisar", 2)
    add_text(doc, "Cuando una planilla tenga errores o inconsistencias, corrige el archivo original y vuelve a cargarlo. Si necesitas quitar una revisión, usa Eliminar de la revisión dentro del detalle del archivo.")

    doc.add_page_break()
    add_heading(doc, "5. Serie A: adjuntar y consolidar", 1)
    add_heading(doc, "Adjuntar y sumar una planilla complementaria", 2)
    add_text(doc, "Para una planilla base de Serie A, usa Adjuntar / sumar o Adjuntar desde el detalle. Selecciona otra planilla .xlsm para asociarla a la base.")
    for item in [
        "La planilla complementaria se analiza antes de ser asociada.",
        "Queda visible en el detalle de la planilla base.",
        "Sus valores suman al resumen y a la descarga consolidada.",
        "Puede quitarse individualmente con el ícono de eliminar.",
    ]:
        add_bullet(doc, item)
    add_callout(doc, "ALCANCE ACTUAL", "La funcionalidad de adjuntar y descargar está habilitada inicialmente solo para Serie A.", fill="FFF8E8", label_color=GOLD_TEXT)
    add_heading(doc, "Descargar un REM consolidado", 2)
    add_text(doc, "En la fila de una planilla Serie A, selecciona Descargar. REMTool genera un archivo .xlsm consolidado utilizando la planilla base y sus complementarias. Mientras se genera el archivo, el botón muestra Generando.... Si el servidor devuelve un error, el panel informa el motivo.")
    add_callout(doc, "RESULTADO", "La descarga combina los datos de la planilla base y de las planillas complementarias asociadas a esa base.")

    add_heading(doc, "6. Exportar el informe de revisión", 1)
    add_text(doc, "Usa Exportar revisión para descargar un informe compatible con Word. El informe incluye el período, el sector seleccionado y, para cada planilla, sus inconsistencias, errores y advertencias. También incorpora los archivos que no pudieron analizarse durante la carga.")
    add_table(doc, ["Incluye", "No reemplaza"], [
        ("Hallazgos por planilla y errores de procesamiento", "La corrección del archivo original."),
        ("Contexto de revisión y sector seleccionado", "Un repositorio central de planillas."),
        ("Evidencia descargable de la revisión", "La validación o aprobación administrativa."),
    ], [4680, 4680])

    add_heading(doc, "7. Cómo interpretar los estados", 1)
    add_table(doc, ["Señal", "Significado", "Acción sugerida"], [
        ("Verde", "Sin errores ni advertencias.", "Puedes continuar con el resumen."),
        ("Amarillo", "Advertencia o exclusión del resumen.", "Abre el detalle y confirma si corresponde corregir o justificar."),
        ("Rojo", "Error de revisión o de procesamiento.", "Corrige la planilla y vuelve a cargarla."),
    ], [1500, 3800, 4060])
    add_heading(doc, "8. Validaciones automáticas", 1)
    for item in [
        "Extensión .xlsm y existencia de la hoja NOMBRE.",
        "Lectura de un mes válido.",
        "Existencia de versión configurada en el sistema.",
        "Identificación del establecimiento por CodDEIS.",
        "Correspondencia entre serie del encabezado y serie de la versión.",
        "Correspondencia con el mes y año seleccionados.",
        "Correspondencia con el establecimiento esperado, cuando aplica.",
        "Reglas de validación configuradas para la versión.",
        "Hojas y prestaciones habilitadas para extraer datos.",
    ]:
        add_bullet(doc, item)
    add_callout(doc, "REGLA DE INCLUSIÓN", "Una diferencia de serie, mes, año o establecimiento se registra como inconsistencia y deja la planilla excluida del resumen. Una regla configurada como advertencia se muestra en amarillo; una regla configurada como error se muestra en rojo.", fill="F4F8FB", label_color=TEAL)

    doc.add_page_break()
    add_heading(doc, "9. Persistencia y limpieza de revisiones", 1)
    add_text(doc, "Las revisiones quedan guardadas en el almacenamiento local del navegador utilizado. Esto permite cambiar de pantalla o volver al Panel REM sin perder las cargas del navegador actual.")
    for item in [
        "Las revisiones no se comparten automáticamente entre computadores, navegadores o perfiles.",
        "Limpiar los datos del navegador puede eliminar las revisiones guardadas.",
        "Limpiar revisiones borra todas las revisiones almacenadas en ese navegador, previa confirmación.",
        "Antes de limpiar, exporta el informe si necesitas conservar evidencia de la revisión.",
    ]:
        add_bullet(doc, item)
    add_callout(doc, "BUENA PRÁCTICA", "Exporta la revisión al cerrar el trabajo del período y conserva el informe junto con los archivos corregidos o aprobados.", fill="EAF7F0", label_color=GREEN_TEXT)

    add_heading(doc, "10. Demostración breve para usuarios", 1)
    add_text(doc, "Este recorrido sirve para una capacitación o una demostración de cinco minutos:")
    for item in [
        "Entra a Panel REM y selecciona un mes, año y sector.",
        "Carga dos o más planillas .xlsm.",
        "Muestra el estado de procesamiento de cada archivo.",
        "Abre una planilla con advertencias o exclusión y explica el panel lateral.",
        "Vuelve al resumen y muestra cómo los puntos se calculan por serie.",
        "En una planilla Serie A, adjunta una complementaria y descarga el REM consolidado.",
        "Finaliza con Exportar revisión para mostrar la evidencia descargable.",
    ]:
        add_number(doc, item)

    add_heading(doc, "11. Observaciones de la revisión funcional", 1)
    for item in [
        "El Panel REM analiza y resume archivos cargados en la sesión del navegador; no es un repositorio central de planillas.",
        "La descarga consolidada y el adjuntar/sumar están implementados inicialmente para Serie A.",
        "El resumen se calcula por serie usando los puntos resumen configurados en el sistema.",
        "La interfaz cuenta con métricas de cobertura en la lógica del panel, pero actualmente esas tarjetas están ocultas; por eso el usuario ve principalmente los puntos resumen y el estado de cada planilla.",
        "El informe de revisión se descarga con extensión .doc compatible con Word, aunque no es un .docx nativo.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "12. Soporte", 1)
    add_text(doc, "Para consultas o sugerencias: ricardocontreras@mpatria.cl", 12, NAVY, True, after=4)
    add_text(doc, "REMTool · DESAM Monte Patria", 10, MUTED, after=0)

    out = OUT / "Manual_Panel_REM.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    build()
